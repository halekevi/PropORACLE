#!/usr/bin/env python3
"""Soccer hit rates and gates workbook.

Reads graded_props from ui_runner/templates (full history) and the latest
soccer step8 board. Writes data/reports/soccer_hit_rates_gates_latest.xlsx.

  py -3.14 scripts/export_soccer_hit_rates_excel.py
  py -3.14 scripts/export_soccer_hit_rates_excel.py --step8 H:\\PropORACLE_main_cp\\outputs\\2026-09-03\\soccer\\step8_soccer_direction_clean.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO / "scripts"))

from rank_best_props_today import _over_d_ok, _under_d_ok  # noqa: E402

_GRADED_DIR = _REPO / "ui_runner" / "templates"
_OUT_XLSX = _REPO / "data" / "reports" / "soccer_hit_rates_gates_latest.xlsx"

TICKET_BANNED = {
    "passes attempted",
    "tackles",
    "fouls",
    "clearances",
    "attempted dribbles",
}
# Sofascore-backed Popular markets (ESPN boxscores lack these columns).
POPULAR = {
    "passes attempted",
    "clearances",
    "attempted dribbles",
}
CORE_TICKET = {
    "shots",
    "shots on target",
    "goalie saves",
    "goals",
    "assists",
    "goal + assist",
    "offsides",
    "cards",
    "fouls drawn",
    "crosses",
    "shots assisted",
}

POSITION_GATES = {
    "saves / goalie saves": "GK only",
    "goals allowed": "GK only",
    "shots / shots on target / goals / assists / key passes": "FWD or MID",
    "tackles": "DEF or MID",
    "clearances": "DEF or GK",
}


def _num(v: object) -> float | None:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    s = str(v).strip()
    if s in ("", "—", "–", "-", "NaN", "nan", "None"):
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _norm_dir(v: object) -> str:
    d = str(v or "").strip().upper()
    if d in ("OVER", "O"):
        return "OVER"
    if d in ("UNDER", "U"):
        return "UNDER"
    return d


def _norm_pick(v: object) -> str:
    p = str(v or "").strip().lower()
    if "goblin" in p:
        return "Goblin"
    if "demon" in p:
        return "Demon"
    if "standard" in p:
        return "Standard"
    return str(v or "").strip() or "Unknown"


def _norm_prop(v: object) -> str:
    s = str(v or "").strip()
    s = re.sub(r"\s+", " ", s)
    if s.lower().endswith(" (combo)"):
        s = s[: -len(" (combo)")].strip()
    return s


def _prop_key(v: object) -> str:
    return _norm_prop(v).casefold()


def _market_group(prop: str) -> str:
    k = _prop_key(prop)
    if k in TICKET_BANNED:
        return "Ticket-banned (noisy)"
    if k in POPULAR:
        return "Popular / Sofascore"
    if k in CORE_TICKET:
        return "Ticket-eligible core"
    return "Other"


def _parse_hit(result: object) -> int | None:
    t = str(result or "").strip().upper()
    if t in ("HIT", "WIN", "W", "1", "TRUE"):
        return 1
    if t in ("MISS", "LOSS", "L", "0", "FALSE"):
        return 0
    return None


def _norm_def(raw: object) -> str:
    s = str(raw or "").strip()
    if not s or s.lower() in {"n/a", "na", "none", "—", "-", "nan", "unknown"}:
        return ""
    low = s.lower()
    if "weak" in low or low in {"bottom", "easy"} or "easy" in low:
        return "Weak"
    if "below" in low:
        return "Below Avg"
    if "elite" in low or low in {"top", "hard", "stingy"} or "hard" in low:
        return "Elite"
    if "above" in low:
        return "Above Avg"
    if low in {"avg", "average"}:
        return "Avg"
    return s


def _d_pass(direction: str, def_tier: str) -> bool:
    if not def_tier:
        return False
    if direction == "OVER":
        return bool(_over_d_ok("SOCCER", def_tier))
    if direction == "UNDER":
        return bool(_under_d_ok("SOCCER", def_tier))
    return False


def _directional_l5(row: dict[str, Any], direction: str) -> int | None:
    lo = _num(row.get("l5_over"))
    if lo is None:
        lo = _num(row.get("last5_over"))
    lu = _num(row.get("l5_under"))
    if lu is None:
        lu = _num(row.get("last5_under"))
    val = lo if direction == "OVER" else lu if direction == "UNDER" else None
    if val is None:
        return None
    if val < 0 or val > 5:
        return None
    return int(round(val))


def load_graded() -> pd.DataFrame:
    paths = sorted(p for p in _GRADED_DIR.glob("graded_props_*.json") if ".bak_" not in p.name)
    rows: list[dict[str, Any]] = []
    for path in paths:
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        file_date = str(raw.get("date") or path.stem.replace("graded_props_", ""))[:10]
        chunk = raw.get("props", raw.get("rows", []))
        if not isinstance(chunk, list):
            continue
        for r in chunk:
            if not isinstance(r, dict):
                continue
            sport = str(r.get("sport") or "").strip().upper()
            if sport not in ("SOCCER", "SOC"):
                continue
            hit = _parse_hit(r.get("result"))
            if hit is None:
                continue
            direction = _norm_dir(r.get("direction") or r.get("over_under"))
            pick = _norm_pick(r.get("pick_type"))
            raw_prop = str(r.get("prop") or r.get("prop_type") or "").strip()
            prop = _norm_prop(raw_prop)
            l5 = _directional_l5(r, direction)
            def_tier = _norm_def(r.get("def_tier") or r.get("DEF_TIER") or r.get("opponent_def_tier"))
            rows.append(
                {
                    "file_date": file_date,
                    "player": str(r.get("player") or "").strip(),
                    "prop": prop or "Unknown",
                    "prop_raw": raw_prop or prop or "Unknown",
                    "prop_key": _prop_key(prop),
                    "market_group": _market_group(prop),
                    "pick_type": pick,
                    "direction": direction,
                    "l5": l5,
                    "hit": int(hit),
                    "def_tier": def_tier or "—",
                    "d_gate": _d_pass(direction, def_tier),
                    "is_combo": "combo" in str(r.get("prop") or "").lower(),
                }
            )
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    df["has_l5"] = df["l5"].notna()
    df["l5_ge4"] = df["l5"].fillna(-1) >= 4
    df["l5_eq5"] = df["l5"].fillna(-1) == 5
    df["l5_eq4"] = df["l5"].fillna(-1) == 4
    df["list_shape"] = (
        df["pick_type"].isin(["Standard", "Goblin"])
        & ~((df["pick_type"] == "Goblin") & (df["direction"] != "OVER"))
    )
    df["list_eligible"] = df["list_shape"] & df["l5_ge4"]
    df["ticket_banned"] = df["prop_key"].isin(TICKET_BANNED)
    df["ticket_shape"] = df["list_shape"] & ~df["ticket_banned"]
    df["ticket_l5"] = df["ticket_shape"] & df["l5_ge4"]
    df["l5_d_ge4"] = df["l5_ge4"] & df["d_gate"]
    df["l5_d_eq5"] = df["l5_eq5"] & df["d_gate"]
    df["l5_4_d"] = df["l5_eq4"] & df["d_gate"]
    return df


def _hr_row(name: str, s: pd.DataFrame, note: str = "") -> dict[str, Any]:
    n = int(len(s))
    hits = int(s["hit"].sum()) if n else 0
    return {
        "pool": name,
        "hits": hits,
        "n": n,
        "misses": n - hits,
        "hit_rate": (hits / n) if n else None,
        "note": note,
    }


def _agg(df: pd.DataFrame, cols: list[str], min_n: int = 1) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    g = df.groupby(cols, dropna=False).agg(hits=("hit", "sum"), n=("hit", "size")).reset_index()
    g["misses"] = g["n"] - g["hits"]
    g["hit_rate"] = (g["hits"] / g["n"]).where(g["n"] > 0)
    g = g[g["n"] >= min_n]
    return g.sort_values(["hit_rate", "n"], ascending=[False, False])


def _find_step8(explicit: str | None) -> Path | None:
    if explicit:
        p = Path(explicit)
        return p if p.exists() else None
    today = datetime.now().strftime("%Y-%m-%d")
    cands = [
        _REPO.parent / "PropORACLE_main_cp" / "outputs" / today / "soccer" / "step8_soccer_direction_clean.xlsx",
        Path(r"H:\PropORACLE_main_cp") / "outputs" / today / "soccer" / "step8_soccer_direction_clean.xlsx",
        _REPO / "outputs" / today / "soccer" / "step8_soccer_direction_clean.xlsx",
        _REPO / "Sports" / "Soccer" / "outputs" / "step8_soccer_direction_clean.xlsx",
        _REPO / "Sports" / "Soccer" / "step8_soccer_direction_clean.xlsx",
    ]
    existing = [p for p in cands if p.exists()]
    if not existing:
        return None
    return max(existing, key=lambda p: p.stat().st_mtime)


def load_live_board(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=0)
    colmap = {str(c).strip().lower(): c for c in df.columns}

    def col(*names: str) -> str | None:
        for n in names:
            if n.lower() in colmap:
                return colmap[n.lower()]
        return None

    c_prop = col("Prop", "prop_type")
    c_pick = col("Pick Type", "pick_type")
    c_dir = col("Direction", "bet_direction")
    c_l5o = col("L5 Over", "last5_over", "l5_over")
    c_l5u = col("L5 Under", "last5_under", "l5_under")
    c_def = col("Def Tier", "DEF_TIER")
    c_player = col("Player")
    c_line = col("Line")
    out = pd.DataFrame()
    out["player"] = df[c_player].astype(str) if c_player else ""
    out["prop"] = df[c_prop].map(_norm_prop) if c_prop else "Unknown"
    out["prop_key"] = out["prop"].map(_prop_key)
    out["market_group"] = out["prop"].map(_market_group)
    out["pick_type"] = df[c_pick].map(_norm_pick) if c_pick else "Unknown"
    out["direction"] = df[c_dir].map(_norm_dir) if c_dir else ""
    out["line"] = pd.to_numeric(df[c_line], errors="coerce") if c_line else np.nan
    out["l5_over"] = pd.to_numeric(df[c_l5o], errors="coerce") if c_l5o else np.nan
    out["l5_under"] = pd.to_numeric(df[c_l5u], errors="coerce") if c_l5u else np.nan
    out["def_tier"] = df[c_def].map(_norm_def) if c_def else ""
    out["def_tier"] = out["def_tier"].replace("", "—")
    l5 = np.where(out["direction"] == "OVER", out["l5_over"], np.where(out["direction"] == "UNDER", out["l5_under"], np.nan))
    out["l5"] = pd.to_numeric(l5, errors="coerce")
    out["has_l5"] = out["l5"].notna()
    out["l5_ge4"] = out["l5"].fillna(-1) >= 4
    out["d_gate"] = [
        _d_pass(d, t if t != "—" else "") for d, t in zip(out["direction"], out["def_tier"])
    ]
    out["list_shape"] = (
        out["pick_type"].isin(["Standard", "Goblin"])
        & ~((out["pick_type"] == "Goblin") & (out["direction"] != "OVER"))
    )
    out["list_eligible"] = out["list_shape"] & out["l5_ge4"]
    out["ticket_banned"] = out["prop_key"].isin(TICKET_BANNED)
    out["gold_shape"] = out["list_eligible"] & out["d_gate"]
    return out


def _style_sheet(ws, df: pd.DataFrame, pct_cols: list[str] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt = PatternFill("solid", fgColor="F2F7FB")
    pct_cols = pct_cols or []
    pct_idx = {i for i, c in enumerate(df.columns, start=1) if c in pct_cols or c == "hit_rate"}
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    for r_i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column), start=2):
        for c_i, cell in enumerate(row, start=1):
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=c_i == 1)
            if r_i % 2 == 0:
                cell.fill = alt
            if c_i in pct_idx and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        sample = [str(col)] + [str(v)[:48] for v in df[col].head(40).tolist()]
        width = min(42, max(12, max(len(x) for x in sample) + 2))
        ws.column_dimensions[letter].width = width


def _write_df(writer: pd.ExcelWriter, name: str, df: pd.DataFrame, pct_cols: list[str] | None = None) -> None:
    if df is None or df.empty:
        pd.DataFrame({"note": ["no rows"]}).to_excel(writer, sheet_name=name, index=False)
        return
    out = df.copy()
    out.to_excel(writer, sheet_name=name, index=False)
    _style_sheet(writer.sheets[name], out, pct_cols=pct_cols)


def build_gates_sheet(date_min: str, date_max: str, graded_n: int, step8_path: str) -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["Best-props list", "Hard list gate", "Directional L5 >= 4 (use last5_over / last5_under when l5_* empty)", "HARD", "No L5 = not listed"],
            ["Best-props list", "Pick types", "Standard Over, Standard Under, Goblin Over. No Demons. Goblin Under dropped.", "HARD", ""],
            ["Best-props list", "Sports on the list", "Always print WNBA / MLB / Soccer / Tennis (thin pools still get a section)", "HARD", ""],
            ["Best-props list", "D (defense)", "Overs vs Weak | Below Avg; Unders vs Elite | Above Avg. Avg never passes D.", "BADGE ONLY", "Failing D still lists, usually Silver"],
            ["Best-props list", "Gold", "L5 + Cover + Delta + Dir + D + Rank all pass (N/A skipped, not a miss)", "BADGE", "0 misses among applicable checks"],
            ["Best-props list", "Silver / Bronze", "Silver = 1 miss; Bronze = 2 misses", "BADGE", ""],
            ["Best-props list", "Platinum / Diamond", "Gold + (L10>=8 and/or season HR>70% min 10). Both = Diamond; one = Platinum", "BADGE", "Soccer saves sit in prop-tier A"],
            ["Combined tickets", "Seed order (all sports)", "L5=5+D → L5=4+D → L5>=4 no-D, then per-sport named-board floors", "HARD", "Stops one sport wiping another"],
            ["Combined tickets", "Noisy soccer legs", "Passes Attempted, Tackles, Fouls, Clearances, Attempted Dribbles excluded", "HARD", "Best-props list still includes them"],
            ["Combined tickets", "Pick hygiene", "No Demon; Goblin Over-only (same as list)", "HARD", "soccer_allowed_leg"],
            ["Combined tickets", "Per-leg hit floor", "2-leg 58% / 3-leg 60% / 4-leg 62% (SOCCER_LEG_MIN_HIT_RATE)", "HARD", "vs global MAIN defaults"],
            ["Combined tickets", "MAIN min leg prob", "Default MAIN_MIN_LEG_PROB = 0.62; directional HR soccer OVER 0.60 / UNDER 0.40", "HARD", "PROP_QUALITY soccer 0.35/0.35"],
            ["Combined tickets", "Standard L5 clear", "Soccer Standard prop×direction ledger clears at L5 >= 4 (not only 5/5)", "HARD", "utils/l5_recency_policy.py"],
            ["Step7 ranker", "Position gates", "Saves/GA = GK; shots/SOT/goals/assists/key passes = FWD/MID; tackles = DEF/MID; clearances = DEF/GK", "VOID", "Wrong position voided, not listed as eligible"],
            ["Step7 ranker", "Starter tier", "STARTER 1.00 / ROTATION 0.92 / SUB 0.75 / UNKNOWN 0.85 rank-score multiplier", "SOFT", ""],
            ["L5 tape", "Popular stats source", "Passes / Clearances / Attempted Dribbles L5 from Sofascore match log (ESPN boxscore lacks those columns)", "DATA", "Shots still ESPN; fantasy dropped in step2"],
            ["This workbook", "Graded window", f"{date_min} to {date_max}  |  decided soccer props = {graded_n:,}", "META", f"Source: {_GRADED_DIR}"],
            ["This workbook", "Live board", step8_path or "(none found)", "META", "Inventory only — not yet graded"],
        ],
        columns=["surface", "gate", "rule", "kind", "notes"],
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--step8", default=None, help="Path to soccer step8 xlsx for live inventory")
    ap.add_argument("--out", default=None, help="Output xlsx path")
    args = ap.parse_args()

    print("Loading graded soccer props…")
    df = load_graded()
    if df.empty:
        print("No graded soccer props found.")
        return 1
    date_min = str(df["file_date"].min())
    date_max = str(df["file_date"].max())
    print(f"  {len(df):,} decided rows  {date_min} to {date_max}")

    std_gob = df[df["pick_type"].isin(["Standard", "Goblin"])]
    list_shape = df[df["list_shape"]]
    with_l5 = list_shape[list_shape["has_l5"]]

    snapshot = pd.DataFrame(
        [
            _hr_row("All decided (any pick)", df, "Includes Demons"),
            _hr_row("Standard + Goblin", std_gob, "Baseline board"),
            _hr_row("Standard only", df[df["pick_type"] == "Standard"]),
            _hr_row("Goblin only", df[df["pick_type"] == "Goblin"]),
            _hr_row("Goblin OVER only", df[(df["pick_type"] == "Goblin") & (df["direction"] == "OVER")], "List/ticket goblin lane"),
            _hr_row("Demon (excluded)", df[df["pick_type"] == "Demon"], "Never listed or ticketed"),
            _hr_row("List shape (Std + Goblin OVER)", list_shape, "Before L5 gate"),
            _hr_row("List shape with L5 tape", with_l5, "Directional last-5 filled"),
            _hr_row("LIST GATE  L5 >= 4", df[df["list_eligible"]], "Hard best-props filter"),
            _hr_row("L5 >= 4 + D pass", df[df["list_eligible"] & df["d_gate"]], "Gold-shaped; D is badge-only on the list"),
            _hr_row("L5 = 5", df[df["list_shape"] & df["l5_eq5"]], "Ticket seed bucket 1 if +D"),
            _hr_row("L5 = 5 + D", df[df["list_shape"] & df["l5_d_eq5"]], "Ticket seed: L5=5+D"),
            _hr_row("L5 = 4 + D", df[df["list_shape"] & df["l5_4_d"]], "Ticket seed: L5=4+D"),
            _hr_row("L5 >= 4 no D", df[df["list_eligible"] & ~df["d_gate"]], "Ticket seed: L5>=4 no-D; list Silver-typical"),
            _hr_row("D pass only (no L5 required)", df[df["list_shape"] & df["d_gate"]], "D alone does not list"),
            _hr_row("Ticket-eligible props, L5>=4", df[df["ticket_l5"]], "Drops passes/tackles/fouls/clearances/dribbles"),
            _hr_row("Ticket-banned props, L5>=4", df[df["list_eligible"] & df["ticket_banned"]], "On the best-props list, off combined tickets"),
            _hr_row("Popular: Passes L5>=4", df[df["list_eligible"] & (df["prop_key"] == "passes attempted")]),
            _hr_row("Popular: Clearances L5>=4", df[df["list_eligible"] & (df["prop_key"] == "clearances")]),
            _hr_row("Popular: Dribbles L5>=4", df[df["list_eligible"] & (df["prop_key"] == "attempted dribbles")]),
            _hr_row("Saves L5>=4", df[df["list_eligible"] & (df["prop_key"] == "goalie saves")]),
            _hr_row("Shots L5>=4", df[df["list_eligible"] & (df["prop_key"] == "shots")]),
            _hr_row("Shots on Target L5>=4", df[df["list_eligible"] & (df["prop_key"] == "shots on target")]),
        ]
    )

    l5_ladder = _agg(with_l5, ["l5"]).sort_values("l5")
    l5_ladder["l5_label"] = l5_ladder["l5"].map(lambda x: f"{int(x)}/5")
    l5_pick = _agg(with_l5, ["pick_type", "direction", "l5"], min_n=15).sort_values(
        ["pick_type", "direction", "l5"]
    )
    l5_d = _agg(with_l5, ["l5", "d_gate"]).sort_values(["l5", "d_gate"], ascending=[True, False])

    by_prop = (
        list_shape.groupby("prop")
        .agg(
            n_all=("hit", "size"),
            hr_all=("hit", "mean"),
            n_with_l5=("has_l5", "sum"),
            n_l5_ge4=("list_eligible", "sum"),
            n_l5_d=("l5_d_ge4", "sum"),
        )
        .reset_index()
    )
    ge4 = list_shape[list_shape["list_eligible"]]
    ge4_d = list_shape[list_shape["list_eligible"] & list_shape["d_gate"]]
    p_ge4 = ge4.groupby("prop").agg(hits_l5_ge4=("hit", "sum"), n_l5_ge4_chk=("hit", "size"), hr_l5_ge4=("hit", "mean"))
    p_d = ge4_d.groupby("prop").agg(hits_l5_d=("hit", "sum"), n_l5_d_chk=("hit", "size"), hr_l5_d=("hit", "mean"))
    by_prop = by_prop.merge(p_ge4, left_on="prop", right_index=True, how="left")
    by_prop = by_prop.merge(p_d, left_on="prop", right_index=True, how="left")
    by_prop["market_group"] = by_prop["prop"].map(_market_group)
    by_prop["ticket_banned"] = by_prop["prop"].map(lambda p: _prop_key(p) in TICKET_BANNED)
    by_prop = by_prop.sort_values(["n_all"], ascending=False)

    prop_dir = _agg(ge4, ["prop", "pick_type", "direction"], min_n=12)
    popular_slice = _agg(
        list_shape[list_shape["prop_key"].isin(POPULAR | TICKET_BANNED | {"goalie saves", "shots", "shots on target"})],
        ["prop", "pick_type", "direction", "l5_ge4"],
        min_n=8,
    )

    group_hr = []
    for gname, sub in list_shape.groupby("market_group"):
        group_hr.append(_hr_row(f"{gname} — all list-shape", sub))
        group_hr.append(_hr_row(f"{gname} — L5>=4", sub[sub["list_eligible"]]))
        group_hr.append(_hr_row(f"{gname} — L5>=4 + D", sub[sub["list_eligible"] & sub["d_gate"]]))
    market_groups = pd.DataFrame(group_hr)

    pos_df = pd.DataFrame(
        [{"prop_family": k, "allowed_positions": v, "action": "void ineligible rows in step7"} for k, v in POSITION_GATES.items()]
    )

    step8 = _find_step8(args.step8)
    live_inv = pd.DataFrame()
    live_by_prop = pd.DataFrame()
    live_list = pd.DataFrame()
    if step8 is not None:
        print(f"Loading live board {step8}")
        live = load_live_board(step8)
        live_inv = (
            live.groupby(["prop", "pick_type", "market_group"], dropna=False)
            .agg(
                n=("player", "size"),
                n_l5=("has_l5", "sum"),
                n_l5_ge4=("l5_ge4", "sum"),
                n_list=("list_eligible", "sum"),
                n_d=("d_gate", "sum"),
                n_list_and_d=("gold_shape", "sum"),
            )
            .reset_index()
            .sort_values(["n"], ascending=False)
        )
        live_inv["l5_fill"] = (live_inv["n_l5"] / live_inv["n"]).where(live_inv["n"] > 0)
        live_inv["list_rate"] = (live_inv["n_list"] / live_inv["n"]).where(live_inv["n"] > 0)
        live_by_prop = (
            live.groupby(["prop", "market_group"], dropna=False)
            .agg(n=("player", "size"), n_l5=("has_l5", "sum"), n_list=("list_eligible", "sum"), n_list_d=("gold_shape", "sum"))
            .reset_index()
            .sort_values("n", ascending=False)
        )
        live_by_prop["l5_fill"] = live_by_prop["n_l5"] / live_by_prop["n"]
        cols = ["player", "prop", "pick_type", "direction", "line", "l5", "def_tier", "d_gate", "ticket_banned", "market_group"]
        live_list = live.loc[live["list_eligible"], cols].sort_values(["prop", "l5", "player"], ascending=[True, False, True])

    gates = build_gates_sheet(date_min, date_max, int(len(df)), str(step8) if step8 else "")

    how = pd.DataFrame(
        [
            ["What this file is", "Current soccer gates (list vs tickets vs ranker) plus graded hit rates. Live 9/3 board is inventory, not results."],
            ["Graded source", str(_GRADED_DIR)],
            ["Window", f"{date_min} to {date_max}"],
            ["L5 definition", "Hits vs today's line over the last 5 games on the bet side (OVER uses last5_over; UNDER uses last5_under)."],
            ["Why baseline HR looks low", "Ungated Standard+Goblin is ~15–20% because Goals/Assists/Demons-adjacent volume misses a lot. The list gate (L5>=4) is the number that matters."],
            ["D is not a list filter", "A Weak/Below Over or Elite/Above Under is Gold-shaped. Avg/blank D still prints if L5>=4 (usually Silver)."],
            ["Ticket vs list", "Passes, tackles, fouls, clearances, dribbles stay on the soccer best-props list and are stripped from combined tickets."],
            ["Popular L5", "Passes / Clearances / Attempted Dribbles come from Sofascore. Graded n is still thinner than Shots/Saves."],
        ],
        columns=["topic", "detail"],
    )

    out_path = Path(args.out) if args.out else _OUT_XLSX
    out_path.parent.mkdir(parents=True, exist_ok=True)
    copies = [out_path]
    dated = _REPO / "outputs" / datetime.now().strftime("%Y-%m-%d") / "soccer_hit_rates_gates.xlsx"
    copies.append(dated)
    if step8 is not None:
        copies.append(step8.parent / "soccer_hit_rates_gates.xlsx")

    first = copies[0]
    print(f"Writing {first}")
    with pd.ExcelWriter(first, engine="openpyxl") as writer:
        _write_df(writer, "How to read", how)
        _write_df(writer, "Gates", gates)
        _write_df(writer, "Snapshot HR", snapshot)
        _write_df(writer, "L5 ladder", l5_ladder)
        _write_df(writer, "L5 x pick x dir", l5_pick)
        _write_df(writer, "L5 x D", l5_d)
        _write_df(writer, "By prop", by_prop)
        _write_df(writer, "Prop x pick x dir L5ge4", prop_dir)
        _write_df(writer, "Market groups", market_groups)
        _write_df(writer, "Popular vs core", popular_slice)
        _write_df(writer, "Position gates", pos_df)
        _write_df(writer, "Live board by prop-pick", live_inv)
        _write_df(writer, "Live board by prop", live_by_prop)
        _write_df(writer, "Live list L5ge4", live_list)

    for extra in copies[1:]:
        extra.parent.mkdir(parents=True, exist_ok=True)
        if extra.resolve() != first.resolve():
            extra.write_bytes(first.read_bytes())
            print(f"Copied {extra}")
    print("Done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
