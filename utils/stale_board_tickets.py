"""Phantom tickets built from a leftover board (not a live PrizePicks slate).

Early-October 2026 NBA1Q slips used the May OKC@SAS Finals step8. Those games
were not on the October schedule, so they show 0 decided. Mark them
``invalid_reason=stale_board`` — not VOID — and keep them out of the forward
counter, grade_history, and P&L shadow.

``--allow-cross-date-fallback`` is capped at ``MAX_CROSS_DATE_FALLBACK_DAYS``.
A March CBB or May WNBA/NBA1Q file is not an overnight board.
"""

from __future__ import annotations

from datetime import date
from typing import Any

INVALID_REASON = "stale_board"
# First date whose NBA1Q tickets came from a real 2026-27 1Q board.
NBA1Q_REAL_BOARD_FROM = "2026-10-20"
MAX_CROSS_DATE_FALLBACK_DAYS = 3
_FINALS_TEAMS = frozenset({"OKC", "SAS"})
_NBA1Q = "NBA1Q"


def days_apart(a: str, b: str) -> int | None:
    """Absolute calendar-day gap, or None if either side is not YYYY-MM-DD."""
    try:
        da = date.fromisoformat(str(a or "")[:10])
        db = date.fromisoformat(str(b or "")[:10])
    except ValueError:
        return None
    return abs((da - db).days)


def cross_date_fallback_allowed(
    chosen: str,
    target: str,
    *,
    max_days: int = MAX_CROSS_DATE_FALLBACK_DAYS,
) -> bool:
    """True only when the fallback date is within ``max_days`` of the slate."""
    gap = days_apart(chosen, target)
    return gap is not None and gap <= int(max_days)


def _norm_team(raw: Any) -> str:
    s = str(raw or "").strip().upper()
    if "@" in s:
        s = s.replace("@", " ")
    s = s.replace("VS", " ").replace(".", " ")
    parts = [p for p in s.split() if p]
    if not parts:
        return ""
    token = parts[-1] if len(parts[-1]) <= 4 else parts[0]
    aliases = {"THUNDER": "OKC", "SPURS": "SAS", "SAN": "SAS", "ANTONIO": "SAS"}
    return aliases.get(token, token)


def _leg_teams(leg: dict[str, Any]) -> set[str]:
    out: set[str] = set()
    for key in ("team", "Team", "opp", "Opp", "opp_team", "opponent", "matchup"):
        tok = _norm_team(leg.get(key))
        if tok:
            out.add(tok)
    matchup = str(leg.get("matchup") or "")
    if "@" in matchup or " vs " in matchup.lower():
        for part in matchup.replace("@", " ").replace("vs", " ").replace("VS", " ").split():
            tok = _norm_team(part)
            if tok:
                out.add(tok)
    return out


def _leg_sport(leg: dict[str, Any], default: str = "") -> str:
    return str(leg.get("sport") or default or "").strip().upper()


def _leg_game_date(leg: dict[str, Any]) -> str:
    for key in ("game_date", "Game Date", "slate_date"):
        v = str(leg.get(key) or "").strip()[:10]
        if len(v) == 10 and v[4] == "-":
            return v
    return ""


def ticket_nba1q_legs(
    ticket: dict[str, Any],
    *,
    group_name: str = "",
) -> list[dict[str, Any]]:
    default = ""
    g = str(group_name or ticket.get("web_group_name") or ticket.get("group_name") or "")
    if "NBA1Q" in g.upper() or "NBA 1Q" in g.upper():
        default = _NBA1Q
    out: list[dict[str, Any]] = []
    for leg in ticket.get("legs") or []:
        if not isinstance(leg, dict):
            continue
        if _leg_sport(leg, default) == _NBA1Q:
            out.append(leg)
    return out


def _okc_sas_pairing(legs: list[dict[str, Any]]) -> bool:
    teams: set[str] = set()
    for leg in legs:
        teams |= _leg_teams(leg)
    return _FINALS_TEAMS <= teams


def _playoff_game_date(gd: str) -> bool:
    return gd.startswith("2026-05") or gd.startswith("2026-06")


def ticket_stale_board_reason(
    ticket: dict[str, Any] | None,
    *,
    slate_date: str = "",
    group_name: str = "",
) -> str | None:
    """Return ``stale_board`` when this slip is a leftover-board phantom, else None.

    Does not mark VOID. Callers set ``invalid`` / ``exclude_from_*``.
    """
    if not isinstance(ticket, dict):
        return None
    if str(ticket.get("invalid_reason") or "").strip().lower() == INVALID_REASON:
        return INVALID_REASON
    ds = str(slate_date or ticket.get("date") or ticket.get("slate_date") or "")[:10]
    gname = str(group_name or ticket.get("web_group_name") or ticket.get("group_name") or "")
    q_legs = ticket_nba1q_legs(ticket, group_name=gname)
    if not q_legs:
        return None
    if ds and ds < NBA1Q_REAL_BOARD_FROM:
        return INVALID_REASON
    gds = {_leg_game_date(leg) for leg in q_legs}
    gds.discard("")
    if any(_playoff_game_date(g) for g in gds):
        return INVALID_REASON
    # Oct 21 still shipped the May OKC@SAS pairing after 10-20's real board.
    # 2026-27 opening week did not host that Finals rematch.
    if ds.startswith("2026-10") and _okc_sas_pairing(q_legs):
        return INVALID_REASON
    return None


def ticket_is_stale_board(
    ticket: dict[str, Any] | None,
    *,
    slate_date: str = "",
    group_name: str = "",
) -> bool:
    return ticket_stale_board_reason(ticket, slate_date=slate_date, group_name=group_name) is not None


def board_row_is_stale_board(row: dict[str, Any] | None, *, slate_date: str = "") -> bool:
    """True for an NBA1Q board row whose game_date is a leftover Finals slate."""
    if not isinstance(row, dict):
        return False
    if str(row.get("sport") or row.get("Sport") or "").strip().upper() != _NBA1Q:
        return False
    ds = str(slate_date or row.get("slate_date") or "")[:10]
    gd = _leg_game_date(row)
    if gd and _playoff_game_date(gd):
        return True
    if ds and ds < NBA1Q_REAL_BOARD_FROM:
        return True
    tid = str(row.get("ticket_id") or "")
    if "2026-05-" in tid or "2026-06-" in tid:
        return True
    team = _norm_team(row.get("team") or row.get("Team"))
    opp = _norm_team(row.get("opp") or row.get("Opp") or row.get("opp_team"))
    # 10-21 re-grade stored one side per row with blank opp (same OKC@SAS leftover).
    if ds.startswith("2026-10") and ds != NBA1Q_REAL_BOARD_FROM:
        if team in _FINALS_TEAMS and (not opp or opp in _FINALS_TEAMS):
            return True
    fake = dict(row)
    fake["legs"] = [row]
    return ticket_is_stale_board(fake, slate_date=ds)


def ticket_leg_row_is_stale_board(row: dict[str, Any] | None) -> bool:
    """Ledger / join-counter filter (early-Oct slate, or Oct OKC@SAS leftover)."""
    if not isinstance(row, dict):
        return False
    if str(row.get("invalid_reason") or "").strip().lower() == INVALID_REASON:
        return True
    sport = str(row.get("sport") or "").strip().upper()
    ds = str(row.get("slate_date") or row.get("date") or "")[:10]
    if sport != _NBA1Q:
        return False
    if ds and ds < NBA1Q_REAL_BOARD_FROM:
        return True
    fake = dict(row)
    fake["sport"] = _NBA1Q
    return board_row_is_stale_board(fake, slate_date=ds)


def annotate_stale_board_ticket(
    ticket: dict[str, Any],
    *,
    slate_date: str = "",
    group_name: str = "",
) -> bool:
    """Stamp invalid/stale_board and exclude flags. Never sets result=VOID."""
    reason = ticket_stale_board_reason(ticket, slate_date=slate_date, group_name=group_name)
    if not reason:
        return False
    ticket["invalid"] = True
    ticket["invalid_reason"] = reason
    ticket["exclude_from_winrate"] = True
    ticket["exclude_from_book"] = True
    ticket["exclude_from_forward"] = True
    # If something previously voided these for 0-decided, undo that label.
    if str(ticket.get("void_reason") or "").strip().lower() in {INVALID_REASON, "stale board"}:
        ticket.pop("void_reason", None)
    res = str(ticket.get("result") or "").strip().upper()
    if res == "VOID":
        ticket["result"] = "INVALID"
    return True


def annotate_ticket_payload(payload: dict[str, Any] | None) -> int:
    """Annotate groups[].tickets and a flat tickets list. Returns stamps applied."""
    if not isinstance(payload, dict):
        return 0
    ds = str(payload.get("date") or payload.get("slate_date") or "")[:10]
    n = 0
    for grp in payload.get("groups") or []:
        if not isinstance(grp, dict):
            continue
        gname = str(grp.get("group_name") or "")
        for t in grp.get("tickets") or []:
            if not isinstance(t, dict):
                continue
            if annotate_stale_board_ticket(t, slate_date=ds, group_name=gname):
                n += 1
            elif _annotate_after_start(t, slate_date=ds, group_name=gname):
                n += 1
    for t in payload.get("tickets") or []:
        if not isinstance(t, dict):
            continue
        if annotate_stale_board_ticket(t, slate_date=ds):
            n += 1
        elif _annotate_after_start(t, slate_date=ds):
            n += 1
    return n


def _annotate_after_start(
    ticket: dict[str, Any], *, slate_date: str = "", group_name: str = ""
) -> bool:
    from utils.published_after_start import annotate_published_after_start_ticket

    return annotate_published_after_start_ticket(
        ticket, slate_date=slate_date, group_name=group_name
    )


def stamp_grade_dataframe(df: Any, *, slate_date: str = "", sport: str = "") -> int:
    """Set invalid_reason=stale_board on leftover-board grade rows. Never VOID."""
    if df is None or getattr(df, "empty", True):
        return 0
    n = 0
    if "invalid_reason" not in df.columns:
        df["invalid_reason"] = ""
    if "exclude_from_forward" not in df.columns:
        df["exclude_from_forward"] = False
    if "exclude_from_winrate" not in df.columns:
        df["exclude_from_winrate"] = False
    sport_u = str(sport or "").strip().upper()
    for idx in df.index:
        rec = df.loc[idx].to_dict()
        if not str(rec.get("sport") or "").strip():
            rec["sport"] = sport_u
        if not board_row_is_stale_board(rec, slate_date=slate_date):
            continue
        df.at[idx, "invalid_reason"] = INVALID_REASON
        df.at[idx, "exclude_from_forward"] = True
        df.at[idx, "exclude_from_winrate"] = True
        res = str(rec.get("result") or "").strip().upper()
        if res == "VOID":
            df.at[idx, "result"] = "INVALID"
        n += 1
    return n


def stamp_graded_xlsx(path: Any, *, slate_date: str = "", sport: str = "NBA1Q") -> int:
    """Patch an existing graded workbook so stale rows carry invalid_reason."""
    from pathlib import Path

    p = Path(path)
    if not p.is_file():
        return 0
    from openpyxl import load_workbook

    wb = load_workbook(p)
    if "Box Raw" not in wb.sheetnames:
        return 0
    ws = wb["Box Raw"]
    headers = [cell.value for cell in ws[1]]
    col_map = {str(h): i + 1 for i, h in enumerate(headers) if h is not None}
    if "invalid_reason" not in col_map:
        idx = len(headers) + 1
        ws.cell(1, idx, "invalid_reason")
        col_map["invalid_reason"] = idx
        headers.append("invalid_reason")
    n = 0
    for r in range(2, int(ws.max_row or 1) + 1):
        rec: dict[str, Any] = {}
        for h, c in col_map.items():
            if h == "invalid_reason":
                continue
            rec[h] = ws.cell(r, c).value
        if not str(rec.get("sport") or rec.get("Sport") or "").strip():
            rec["sport"] = sport
        if rec.get("opp") is None and rec.get("opp_team") is not None:
            rec["opp"] = rec.get("opp_team")
        if not board_row_is_stale_board(rec, slate_date=slate_date):
            continue
        ws.cell(r, col_map["invalid_reason"], INVALID_REASON)
        if "result" in col_map:
            res = str(ws.cell(r, col_map["result"]).value or "").strip().upper()
            if res == "VOID":
                ws.cell(r, col_map["result"], "INVALID")
        n += 1
    if n:
        wb.save(p)
    return n
