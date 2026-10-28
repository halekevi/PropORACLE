"""Machine-hop Parquet sidecar next to CSV/XLSX."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from proporacle.data.table_io import (
    copy_parquet_sidecar,
    format_step8_meta_banner,
    infer_slate_date_from_path,
    nba_required_columns_all_blank,
    parquet_available,
    parquet_sidecar_path,
    read_step8_meta_freshness,
    read_table,
    read_table_str,
    step8_freshness_label,
    table_exists,
    write_excel_sheets,
    write_parquet_sidecar,
)


pytestmark = pytest.mark.skipif(not parquet_available(), reason="pyarrow/fastparquet not installed")


def test_parquet_sidecar_preferred_over_csv(tmp_path: Path):
    csv_path = tmp_path / "step8_wnba_direction.csv"
    pd.DataFrame({"player": ["A"], "line": [1.5]}).to_csv(csv_path, index=False)
    write_parquet_sidecar(pd.DataFrame({"player": ["B"], "line": [9.5]}), csv_path)
    assert parquet_sidecar_path(csv_path).is_file()
    df = read_table(csv_path)
    assert list(df["player"]) == ["B"]
    assert float(df["line"].iloc[0]) == 9.5


def test_read_falls_back_to_csv_when_parquet_disabled(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    csv_path = tmp_path / "board.csv"
    pd.DataFrame({"player": ["csv"]}).to_csv(csv_path, index=False)
    write_parquet_sidecar(pd.DataFrame({"player": ["pq"]}), csv_path)
    monkeypatch.setattr("proporacle.data.table_io.parquet_available", lambda: False)
    df = read_table(csv_path)
    assert list(df["player"]) == ["csv"]


def test_table_exists_sees_sidecar_only(tmp_path: Path):
    xlsx = tmp_path / "step8_mlb_direction_clean.xlsx"
    write_parquet_sidecar(pd.DataFrame({"player": ["X"]}), xlsx)
    assert not xlsx.is_file()
    assert table_exists(xlsx)
    df = read_table(xlsx)
    assert list(df["player"]) == ["X"]


def test_copy_parquet_sidecar(tmp_path: Path):
    src = tmp_path / "step8.xlsx"
    dest = tmp_path / "dated" / "step8_clean.xlsx"
    write_parquet_sidecar(pd.DataFrame({"n": [1]}), src)
    copied = copy_parquet_sidecar(src, dest)
    assert copied is not None and copied.is_file()
    assert list(read_table(dest)["n"]) == [1]


def test_read_table_str_stringifies_parquet(tmp_path: Path):
    xlsx = tmp_path / "step7_ranked.xlsx"
    write_parquet_sidecar(pd.DataFrame({"player": ["A"], "line": [12.5], "n": [None]}), xlsx)
    df = read_table_str(xlsx, sheet="ALL")
    assert list(df["player"]) == ["A"]
    assert df["line"].iloc[0] == "12.5"
    assert df["n"].iloc[0] == ""


def test_write_excel_sheets_roundtrip(tmp_path: Path):
    path = tmp_path / "step7_ranked.xlsx"
    write_excel_sheets(
        path,
        {
            "ALL": pd.DataFrame({"Tier": ["A"], "player": ["X"]}),
            "ELIGIBLE": pd.DataFrame({"Tier": ["A"], "player": ["X"]}),
        },
        stamp_meta=False,
    )
    assert path.is_file()
    df = pd.read_excel(path, sheet_name="ALL")
    assert list(df["player"]) == ["X"]
    assert list(df["Tier"]) == ["A"]


def test_step8_meta_banner_marks_stale(tmp_path: Path):
    assert step8_freshness_label("2026-10-20", "2026-10-21") == "STALE"
    assert step8_freshness_label("2026-10-21", "2026-10-21") == "OK"
    banner, fresh = format_step8_meta_banner(
        slate_date="2026-10-20", run_date="2026-10-21", build_time="2026-10-20 09:00 ET"
    )
    assert fresh == "STALE"
    assert "STALE" in banner
    assert "slate_date=2026-10-20" in banner

    dated = tmp_path / "outputs" / "2026-10-20" / "wnba" / "step8_wnba_direction_clean.xlsx"
    dated.parent.mkdir(parents=True)
    write_excel_sheets(
        dated,
        {"ALL": pd.DataFrame({"Tier": ["A"], "player": ["Chen"]})},
        run_date="2026-10-21",
    )
    assert infer_slate_date_from_path(dated) == "2026-10-20"
    meta = pd.read_excel(dated, sheet_name="_META")
    assert list(meta["freshness"]) == ["STALE"]
    assert list(meta["slate_date"]) == ["2026-10-20"]
    assert "STALE" in str(meta["banner"].iloc[0])
    # Banner lives on _META only. Naive pd.read_excel must see real headers.
    raw = pd.read_excel(dated, sheet_name="ALL")
    assert list(raw["player"]) == ["Chen"]
    assert list(read_table(dated, prefer_parquet=False)["player"]) == ["Chen"]


def test_nba_blank_safety_columns_mark_freshness_fail(tmp_path: Path):
    assert step8_freshness_label("2026-10-28", "2026-10-28") == "OK"
    assert step8_freshness_label("2026-10-28", "2026-10-28", column_fail=True) == "FAIL"
    assert step8_freshness_label("2026-10-20", "2026-10-21", column_fail=True) == "STALE"

    blank = pd.DataFrame(
        {
            "Player": ["X"],
            "L5 Current RS N": [None],
            "Current Season RS N": [None],
            "Clean Minutes": [None],
            "Minutes Flag": [None],
        }
    )
    assert nba_required_columns_all_blank(blank) == [
        "L5 Current RS N",
        "Current Season RS N",
        "Clean Minutes",
        "Minutes Flag",
    ]
    filled = blank.copy()
    filled["L5 Current RS N"] = [3]
    filled["Current Season RS N"] = [4]
    filled["Clean Minutes"] = [True]
    filled["Minutes Flag"] = ["ok"]
    assert nba_required_columns_all_blank(filled) == []

    dated = tmp_path / "outputs" / "2026-10-28" / "nba" / "step8_all_direction_clean.xlsx"
    dated.parent.mkdir(parents=True)
    write_excel_sheets(dated, {"ALL": blank}, run_date="2026-10-28")
    assert read_step8_meta_freshness(dated) == "FAIL"
    meta = pd.read_excel(dated, sheet_name="_META")
    assert "FAIL" in str(meta["banner"].iloc[0])
    assert "L5 Current RS N" in str(meta["blank_cols"].iloc[0])

    nfl = tmp_path / "outputs" / "2026-10-28" / "nfl" / "step8_nfl_direction_clean.xlsx"
    nfl.parent.mkdir(parents=True)
    write_excel_sheets(nfl, {"ALL": blank}, run_date="2026-10-28")
    assert read_step8_meta_freshness(nfl) == "OK"

    wnba = tmp_path / "outputs" / "2026-10-28" / "wnba" / "step8_wnba_direction_clean.xlsx"
    wnba.parent.mkdir(parents=True)
    write_excel_sheets(wnba, {"ALL": blank}, run_date="2026-10-28")
    assert read_step8_meta_freshness(wnba) == "OK"


def test_read_table_skips_legacy_row0_banner(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from openpyxl import Workbook

    from proporacle.data.table_io import repair_banner_headers

    monkeypatch.setattr("proporacle.data.table_io.parquet_available", lambda: False)
    dated = tmp_path / "step8_nba_direction_clean.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "ALL"
    ws["A1"] = "slate_date=2026-10-23 | build_time=x | run_date=2026-10-23"
    ws["A2"] = "player"
    ws["B2"] = "line"
    ws["A3"] = "SGA"
    ws["B3"] = 28.5
    wb.save(dated)
    df = read_table(dated, prefer_parquet=False)
    assert list(df["player"]) == ["SGA"]
    broken = pd.read_excel(dated, sheet_name="ALL")
    assert any(str(c).startswith("slate_date=") for c in broken.columns)
    assert list(repair_banner_headers(broken)["player"]) == ["SGA"]


def test_parquet_mixed_na_rank_and_numeric_assign(tmp_path: Path):
    xlsx = tmp_path / "step8_tennis_direction_clean.xlsx"
    df = pd.DataFrame(
        {
            "player": ["A", "B", "C"],
            "OVERALL_DEF_RANK": [12, "N/A", 3],
            "l5_over": ["4", "5", None],
        }
    )
    assert write_parquet_sidecar(df, xlsx) is not None
    got = read_table(xlsx)
    got.loc[got.index, "l5_over"] = [4.0, 5.0, 6.0]
    assert list(got["l5_over"]) == [4.0, 5.0, 6.0]
    assert "N/A" in list(got["OVERALL_DEF_RANK"].astype(str))


def test_existing_arrow_string_l5_over_accepts_float(tmp_path: Path):
    """Boards written before the numpy hop still assign L5 counts in combined."""
    xlsx = tmp_path / "step8_mlb_direction_clean.xlsx"
    raw = pd.DataFrame({"player": ["A", "B"], "l5_over": ["4", "5"]})
    pq = parquet_sidecar_path(xlsx)
    raw.convert_dtypes().to_parquet(pq, index=False)
    got = read_table(xlsx)
    got.loc[got.index, "l5_over"] = [4.0, 5.0]
    assert list(got["l5_over"]) == [4.0, 5.0]


def test_xlsx_unkeyed_sidecar_is_not_used_for_all_sheet(tmp_path: Path):
    """Legacy stem.parquet may be ELIGIBLE. ALL must come from Excel or stem_ALL."""
    xlsx = tmp_path / "step6_ranked_cbb.xlsx"
    write_excel_sheets(
        xlsx,
        {
            "ALL": pd.DataFrame(
                {"player": ["G", "D", "S"], "pick_type": ["Goblin", "Demon", "Standard"]}
            ),
            "ELIGIBLE": pd.DataFrame({"player": ["S"], "pick_type": ["Standard"]}),
        },
        stamp_meta=False,
    )
    eligible = pd.DataFrame({"player": ["S"], "pick_type": ["Standard"]})
    eligible.to_parquet(xlsx.with_suffix(".parquet"), index=False)
    got = read_table(xlsx, sheet="ALL")
    assert len(got) == 3
    assert set(got["pick_type"].astype(str)) == {"Goblin", "Demon", "Standard"}


def test_xlsx_keyed_all_sidecar_is_used(tmp_path: Path):
    xlsx = tmp_path / "step8_tennis_direction_clean.xlsx"
    write_excel_sheets(
        xlsx,
        {"ALL": pd.DataFrame({"player": ["excel"], "pick_type": ["Standard"]})},
        stamp_meta=False,
    )
    write_parquet_sidecar(
        pd.DataFrame({"player": ["sidecar"], "pick_type": ["Goblin"]}),
        xlsx,
        sheet="ALL",
    )
    got = read_table(xlsx, sheet="ALL")
    assert list(got["player"]) == ["sidecar"]
    assert list(got["pick_type"]) == ["Goblin"]


def test_step7b_refresh_keeps_parquet_in_sync(tmp_path: Path):
    """step8 prefers parquet; scoring the xlsx must rewrite the sidecar too."""
    xlsx = tmp_path / "step7_ranked.xlsx"
    df = pd.DataFrame({"player": ["A"], "eligible": [1], "ml_prob": [0.4]})
    write_excel_sheets(xlsx, {"ALL": df, "ELIGIBLE": df})
    write_parquet_sidecar(df, xlsx)
    scored = df.copy()
    scored["ml_prob"] = 0.81
    write_excel_sheets(xlsx, {"ALL": scored, "ELIGIBLE": scored})
    write_parquet_sidecar(scored, xlsx)
    hop = read_table(xlsx)
    assert float(hop["ml_prob"].iloc[0]) == pytest.approx(0.81)
